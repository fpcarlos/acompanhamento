import csv
from openpyxl import load_workbook
from io import TextIOWrapper
import unicodedata
import re

REQUIRED_COLUMNS = [
    'Código da Ação',
    'Área de Atuação',
    'Descrição da Ação',
    'Unidade Técnica Responsável',
    'Tipo de Produto / Instrumento',
    'Indicador',
    'Meta',
    'Ano',
    'Status',
]


def normalize_str(s: str) -> str:
    if s is None:
        return ''
    s = str(s)
    # remove accents
    s = unicodedata.normalize('NFKD', s)
    s = ''.join([c for c in s if not unicodedata.combining(c)])
    s = s.lower()
    # remove non-alphanumeric
    s = re.sub(r'[^a-z0-9]', '', s)
    return s


CANONICAL_COLUMNS = {normalize_str(c): c for c in REQUIRED_COLUMNS}


def map_headers(headers):
    """Map given headers to canonical expected header names.

    Returns (mapping, missing) where mapping is dict canonical -> actual header name and
    missing is list of canonical names missing.
    """
    mapping = {}
    norm_to_actual = {normalize_str(h): h for h in headers}
    for canon_norm, canon_display in CANONICAL_COLUMNS.items():
        if canon_norm in norm_to_actual:
            mapping[canon_display] = norm_to_actual[canon_norm]
        else:
            # try partial match: any header whose normalized contains canon_norm
            found = None
            for nh, actual in norm_to_actual.items():
                if canon_norm in nh:
                    found = actual
                    break
            if found:
                mapping[canon_display] = found
    missing = [c for c in REQUIRED_COLUMNS if c not in mapping]
    return mapping, missing


def read_table_from_file(file_obj):
    """Read CSV or XLSX into list of dicts (header row expected).

    Returns (rows, errors) where rows is a list of dict and errors is a list of strings.
    """
    fname = getattr(file_obj, 'name', '')
    if fname.lower().endswith('.csv'):
        # Decode bytes to text
        text = TextIOWrapper(file_obj.file, encoding='utf-8') if hasattr(file_obj, 'file') else TextIOWrapper(file_obj, encoding='utf-8')
        reader = csv.DictReader(text)
        rows = list(reader)
        return rows, []

    if fname.lower().endswith('.xlsx'):
        wb = load_workbook(file_obj, read_only=True)
        ws = wb.active
        rows = []
        headers = [cell.value for cell in next(ws.rows)]
        for r in ws.iter_rows(min_row=2, values_only=True):
            row = {headers[i]: r[i] for i in range(len(headers))}
            rows.append(row)
        return rows, []

    return [], [f'Unsupported file type: {fname}']


def validate_columns(headers):
    # Use mapping to allow flexible header names and order
    mapping, missing = map_headers(list(headers))
    return missing

